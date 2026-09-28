"""
Genera el manual del mod: docs/manual/MANUAL.md -> docs/manual/FemClothes_Manual.docx

Uso:  python tools/generar_manual.py [--destino <carpeta>]

1. Arma las imágenes en docs/manual/img/ (íconos de ítems, recortes de
   capturas del juego y vistas previas de redes/arneses/motivos generadas
   con los MISMOS algoritmos que el mod — ver PatronGenerador y
   ClothingTextureCache). Si cambia un algoritmo en Java, hay que
   reflejarlo acá para que el manual no mienta.
2. Convierte el Markdown (subconjunto: títulos, párrafos, listas, tablas,
   imágenes, citas, **negrita**, *itálica*, `código`) a Word.
3. Si se pasa --destino (ej. la carpeta de Google Drive), copia ahí el
   .docx — así se sube solo después de cada actualización.

Pedido original (2026-09-28): "un documento que tenga toda la descripcion
de la funcionalidad del mod... que lo actualicemos y lo subamos a drive
despues de cada cambio".
"""
import argparse
import datetime
import math
import os
import re
import shutil

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(RAIZ, 'docs', 'manual')
IMG = os.path.join(DOCS, 'img')
TEX_ITEM = os.path.join(RAIZ, 'src', 'main', 'resources', 'assets', 'femclothes', 'textures', 'item')
SALIDA = os.path.join(DOCS, 'FemClothes_Manual.docx')

# Capturas elegidas: (archivo de salida, ruta de origen, recorte (izq, arriba, der, abajo) o None)
MODRINTH = r'C:\Users\feder\AppData\Roaming\ModrinthApp\profiles\Fabric 1.21.1 (2)\screenshots'
SHAREX = r'C:\Users\feder\OneDrive\Documentos\ShareX\Screenshots\2026-09'
CAPTURAS = [
    ('captura_jugador_1.png', os.path.join(MODRINTH, '2026-09-26_22.18.15.png'), None),
    ('captura_jugador_2.png', os.path.join(SHAREX, 'java_aEgZ9lKuZm.png'), 'ventana'),
    ('captura_modelado_pantalon.png', os.path.join(MODRINTH, '2026-09-27_01.50.03.png'), None),
    ('captura_modelado_medias.png', os.path.join(SHAREX, 'javaw_C5qQADD2wd.png'), 'ventana'),
]

PIEL = (224, 180, 150)
TELA = (214, 96, 150)
FONDO = (46, 38, 42)
METAL = (200, 204, 212)


# ── imágenes ──────────────────────────────────────────────────────────

def capturas():
    for nombre, origen, recorte in CAPTURAS:
        if not os.path.exists(origen):
            print(f'  (falta la captura {origen}, se saltea)')
            continue
        im = Image.open(origen).convert('RGB')
        if recorte == 'ventana':
            # Captura de pantalla completa de ShareX: sin la barra de título
            # de la ventana ni la barra de tareas de Windows.
            w, h = im.size
            im = im.crop((0, round(h * 0.03), w, round(h * 0.955)))
        im.thumbnail((1600, 900))
        im.save(os.path.join(IMG, nombre))


def iconos():
    """Grilla con los íconos de los ítems principales, ampliados sin suavizar."""
    nombres = [
        'corte_normal_corta_redondo', 'socks_solid', 'pantalon', 'calientabrazos', 'pollera',
        'molde_rango_medio', 'molde_torso_medio', 'molde_cuello_redondo', 'molde_calce_normal', 'molde_de_corte',
        'molde_red_fina', 'molde_red_hexagonal', 'molde_red_encaje', 'molde_red_arnes_x', 'molde_red_lisa',
        'pattern_stripe_alt', 'pattern_corazones', 'pattern_estrellas', 'pattern_lunares', 'pattern_vichy',
    ]
    celdas = []
    for n in nombres:
        p = os.path.join(TEX_ITEM, n + '.png')
        if os.path.exists(p):
            celdas.append(Image.open(p).convert('RGBA').resize((64, 64), Image.NEAREST))
    cols = 5
    filas = math.ceil(len(celdas) / cols)
    hoja = Image.new('RGBA', (cols * 80, filas * 80), (243, 232, 208, 255))
    for i, c in enumerate(celdas):
        hoja.paste(c, ((i % cols) * 80 + 8, (i // cols) * 80 + 8), c)
    hoja.save(os.path.join(IMG, 'iconos.png'))


def _dist_seg(px, py, ax, ay, bx, by):
    dx, dy = bx - ax, by - ay
    l2 = dx * dx + dy * dy
    t = 0 if l2 == 0 else max(0, min(1, ((px - ax) * dx + (py - ay) * dy) / l2))
    return math.hypot(px - (ax + t * dx), py - (ay + t * dy))


R3 = math.sqrt(3)


def _hilo(tipo, x, y):
    """Misma lógica que ClothingTextureCache#esHilo (True = tela)."""
    if tipo == 'fina':
        return (x + y) % 4 < 1 or (x - y) % 4 < 1
    if tipo == 'gruesa':
        return (x + y) % 9 < 2 or (x - y) % 9 < 2
    if tipo == 'hexagonal':
        r, h = 8, 2
        px, py = x + .5, y + .5
        af = r * R3
        fc = math.floor(py / af + .5)
        mejor = 1e9
        for f in range(fc - 1, fc + 2):
            cy = f * af
            off = r if f % 2 else 0
            cc = math.floor((px - off) / (2 * r) + .5)
            for c in range(cc - 1, cc + 2):
                cx = c * 2 * r + off
                dx, dy = abs(px - cx), abs(py - cy)
                mejor = min(mejor, max(dx, dx / 2 + dy * R3 / 2))
        return mejor >= r - h
    if tipo == 'perforada':
        paso, rad = 10, 3
        px, py = x + .5, y + .5
        f = math.floor(py / paso)
        off = paso / 2 if f % 2 else 0
        cx = math.floor((px - off) / paso) * paso + off + paso / 2
        cy = f * paso + paso / 2
        return (px - cx) ** 2 + (py - cy) ** 2 > (rad + .5) ** 2
    if tipo == 'encaje':
        c, h = 16, 2
        u, v = (x + y) % c, (x - y) % c
        if u < h or v < h:
            return True
        m = c / 2 + h / 2
        return (u - m) ** 2 + (v - m) ** 2 <= 8
    if tipo == 'rayas':
        if y % 6 >= 2:
            return True
        return (x + ((y // 6) % 2) * 4) % 8 < 2
    if tipo == 'escocesa':
        return x % 8 < 2 or y % 8 < 2
    return True


def redes():
    tipos = ['fina', 'gruesa', 'hexagonal', 'perforada', 'encaje', 'rayas', 'escocesa']
    W, H, B = 32, 80, 4
    hoja = Image.new('RGB', (len(tipos) * (W + 10) + 10, H + 20), FONDO)
    for i, t in enumerate(tipos):
        for y in range(H):
            for x in range(W):
                borde = y < B or y >= H - B
                hoja.putpixel((10 + i * (W + 10) + x, 10 + y), TELA if borde or _hilo(t, x, y) else PIEL)
    hoja = hoja.resize((hoja.width * 4, hoja.height * 4), Image.NEAREST)
    hoja.save(os.path.join(IMG, 'redes.png'))


def arneses():
    """Espalda (arriba, dada vuelta como se ve al doblar por el hombro), tapa de hombros y frente del torso."""
    A1, A2, T = 0.28, 0.72, 4
    M = T / 2

    def anillos(k, w, h):
        if k == 'x':
            return [(w / 2, h * (A2 - A1) / (2 * (1 - A1)))]
        if k == 't':
            return [(w * A1, h / 2), (w * A2, h / 2)]
        return [(w / 2, h / 3), (w / 2, 2 * h / 3)]

    def tira(k, x, y, w, h):
        if k == 'x':
            return _dist_seg(x, y, A1 * w, 0, w, h) < M or _dist_seg(x, y, A2 * w, 0, 0, h) < M
        if k == 't':
            return abs(x - w * A1) < M or abs(x - w * A2) < M or abs(y - h / 2) < M
        return abs(y - h / 3) < M or abs(y - 2 * h / 3) < M or ((abs(x - w * A1) < M or abs(x - w * A2) < M) and y < h / 3)

    def cara(k, w, h):
        im = Image.new('RGB', (w, h))
        for y in range(h):
            for x in range(w):
                cx, cy = x + .5, y + .5
                c = TELA if (cy > h - 4 or tira(k, cx, cy, w, h)) else PIEL
                for ax, ay in anillos(k, w, h):
                    d = math.hypot(cx - ax, cy - ay)
                    if d <= 2.2:
                        c = PIEL
                    elif d <= 5.2:
                        f = 1 + ((ax - cx) + (ay - cy)) / 5.2 * 0.25
                        c = tuple(min(255, int(v * f)) for v in METAL)
                im.putpixel((x, y), c)
        return im

    def tapa(w, d):
        im = Image.new('RGB', (w, d))
        for y in range(d):
            for x in range(w):
                im.putpixel((x, y), TELA if abs(x + .5 - w * A1) < M or abs(x + .5 - w * A2) < M else PIEL)
        return im

    W, H, D = 64, 72, 32
    hoja = Image.new('RGB', (3 * (W + 10) + 10, D + 2 * H + 20), FONDO)
    for i, k in enumerate('xtb'):
        ox = 10 + i * (W + 10)
        hoja.paste(cara(k, W, H).transpose(Image.ROTATE_180), (ox, 10))
        hoja.paste(tapa(W, D), (ox, 10 + H))
        hoja.paste(cara(k, W, H), (ox, 10 + H + D))
    hoja = hoja.resize((hoja.width * 3, hoja.height * 3), Image.NEAREST)
    hoja.save(os.path.join(IMG, 'arneses.png'))


CORAZON = [".XX...XX.", "XXXX.XXXX", "XXXXXXXXX", "XXXXXXXXX", ".XXXXXXX.", "..XXXXX..", "...XXX...", "....X...."]


def _pinta(s, sx, sy):
    return 0 <= sy < len(s) and 0 <= sx < len(s[0]) and s[sy][sx] == 'X'


def _hash(a, b, s):
    m = (1 << 64) - 1
    h = (a * 0x9E3779B97F4A7C15 + b * 0xC2B2AE3D27D4EB4F + s * 0x165667B19E3779F9) & m
    h ^= h >> 33
    h = (h * 0xFF51AFD7ED558CCD) & m
    h ^= h >> 33
    return h


def _rol_motivo(rep, u, v, w, h, k=2, semilla=0):
    """Misma lógica que PatronGenerador#rolSprite: (rol, n° de repetición) — rol 0 nada, 1 relleno, 2 contorno."""
    m = CORAZON
    A, L = len(m[0]), len(m)
    an, al = A * k, L * k
    kk = k
    n = 0
    if rep == 'u':
        kk = k * 2.5
        lu, lv = u - w * 0.375, v - 0.5 * h
    else:
        cols = max(1, round(w / (max(an, al) * 1.8)))
        cx, cy = w / cols, max(an, al) * 1.8
        f = math.floor(v / cy)
        uu = u + (cx / 2 if rep == 'l' and f % 2 else 0)
        col = math.floor(uu / cx) % cols
        n = col + f
        lu = uu - math.floor(uu / cx) * cx - cx / 2
        lv = v - f * cy - cy / 2
        if rep == 'd':
            hh = _hash(col, f, semilla)
            if (hh & 0xFF) / 255 < 0.3:
                return 0, n
            lu -= (((hh >> 8) & 0xFF) / 255 * 2 - 1) * max(0, (cx - an) / 2 - k)
            lv -= (((hh >> 16) & 0xFF) / 255 * 2 - 1) * max(0, (cy - al) / 2 - k)
    sx, sy = math.floor(lu / kk + A / 2), math.floor(lv / kk + L / 2)
    if _pinta(m, sx, sy):
        return 1, n
    if any(_pinta(m, sx + dx, sy + dy) for dx in (-1, 0, 1) for dy in (-1, 0, 1)):
        return 2, n
    return 0, n


def motivos():
    """Corazones en las 4 repeticiones sobre una pierna desplegada (4 caras), con contorno y Alternar en la última fila."""
    W, H = 128, 96
    rojo, rosa, negro, base = (200, 40, 90), (240, 150, 180), (40, 30, 35), (245, 232, 236)
    hoja = Image.new('RGB', (4 * (W + 10) + 10, 2 * (H + 10) + 10), FONDO)
    for fila in range(2):
        for i, rep in enumerate('glud'):
            for y in range(H):
                for x in range(W):
                    rol, n = _rol_motivo(rep, x + .5, y + .5, W, H)
                    if fila == 0:
                        c = rojo if rol == 1 else base
                    else:
                        # Contorno sí + Alternar entre rojo y rosa.
                        c = negro if rol == 2 else ((rojo, rosa)[n % 2] if rol == 1 else base)
                    if x % 32 == 0:
                        c = tuple(int(v * 0.85) for v in c)
                    hoja.putpixel((10 + i * (W + 10) + x, 10 + fila * (H + 10) + y), c)
    hoja = hoja.resize((hoja.width * 2, hoja.height * 2), Image.NEAREST)
    hoja.save(os.path.join(IMG, 'motivos.png'))


# ── Markdown -> docx ─────────────────────────────────────────────────

INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')


def _inline(parrafo, texto, negrita=False):
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith('**'):
            r = parrafo.add_run(trozo[2:-2])
            r.bold = True
        elif trozo.startswith('`'):
            r = parrafo.add_run(trozo[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x6B, 0x2E, 0x4E)
        elif trozo.startswith('*'):
            r = parrafo.add_run(trozo[1:-1])
            r.italic = True
        else:
            r = parrafo.add_run(trozo)
        if negrita:
            r.bold = True


def _tabla(doc, filas):
    celdas = [[c.strip() for c in f.strip().strip('|').split('|')] for f in filas]
    celdas = [f for f in celdas if not all(re.fullmatch(r':?-{2,}:?', c) for c in f)]
    t = doc.add_table(rows=len(celdas), cols=len(celdas[0]))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(celdas):
        for j, c in enumerate(fila[:len(celdas[0])]):
            p = t.cell(i, j).paragraphs[0]
            _inline(p, c, negrita=(i == 0))
            for r in p.runs:
                r.font.size = Pt(9.5)
    doc.add_paragraph()


def convertir():
    with open(os.path.join(DOCS, 'MANUAL.md'), encoding='utf-8') as f:
        lineas = f.read().replace('{{FECHA}}', datetime.date.today().strftime('%d/%m/%Y')).split('\n')

    doc = Document()
    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(11)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.2)

    i = 0
    primer_titulo = True
    while i < len(lineas):
        l = lineas[i]
        if not l.strip():
            i += 1
            continue
        if l.startswith('|'):
            bloque = []
            while i < len(lineas) and lineas[i].startswith('|'):
                bloque.append(lineas[i])
                i += 1
            _tabla(doc, bloque)
            continue
        m = re.match(r'!\[(.*?)\]\((.*?)\)', l.strip())
        if m:
            ruta = os.path.join(DOCS, m.group(2))
            if os.path.exists(ruta):
                ancho = Image.open(ruta).size[0]
                doc.add_picture(ruta, width=Cm(16 if ancho > 400 else 10))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(m.group(1))
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            i += 1
            continue
        if l.startswith('# '):
            if primer_titulo:
                doc.add_heading(l[2:], level=0)
                primer_titulo = False
            else:
                doc.add_page_break()
                doc.add_heading(l[2:], level=1)
        elif l.startswith('## '):
            doc.add_heading(l[3:], level=1)
        elif l.startswith('### '):
            doc.add_heading(l[4:], level=2)
        elif l.startswith('> '):
            p = doc.add_paragraph()
            _inline(p, l[2:])
            for r in p.runs:
                r.italic = True
        elif re.match(r'\s*- ', l):
            nivel = (len(l) - len(l.lstrip())) // 2
            p = doc.add_paragraph(style='List Bullet 2' if nivel else 'List Bullet')
            _inline(p, l.strip()[2:])
        elif re.match(r'\s*\d+\. ', l):
            p = doc.add_paragraph(style='List Number')
            _inline(p, re.sub(r'^\s*\d+\. ', '', l))
        else:
            p = doc.add_paragraph()
            _inline(p, l.strip())
        i += 1

    doc.save(SALIDA)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--destino', help='carpeta a la que copiar el .docx (ej. Google Drive)')
    args = ap.parse_args()
    os.makedirs(IMG, exist_ok=True)
    print('imágenes...')
    capturas()
    iconos()
    redes()
    arneses()
    motivos()
    print('docx...')
    convertir()
    print('listo:', SALIDA)
    if args.destino:
        os.makedirs(args.destino, exist_ok=True)
        shutil.copy2(SALIDA, os.path.join(args.destino, os.path.basename(SALIDA)))
        print('copiado a', args.destino)


if __name__ == '__main__':
    main()
